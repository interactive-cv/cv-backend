# Тесты извлечения текста из файлов ТЗ.
import io

import pytest

from app.services.spec_extractor import extract_spec


def _make_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("ТЗ", 0)
    d.add_paragraph("Требование один")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_extract_success():
    text, elements, file_type = extract_spec("spec.docx", _make_docx())
    assert file_type == "docx"
    assert "Требование один" in text
    assert elements >= 2


def test_docx_old_word_format_friendly_error():
    """Старый .doc (OLE2) с расширением .docx — понятная ошибка, не 500."""
    old_doc = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    with pytest.raises(ValueError, match="старый формат .doc"):
        extract_spec("spec.docx", old_doc)


def test_docx_rtf_friendly_error():
    rtf = b"{\\rtf1\\ansi some requirement}"
    with pytest.raises(ValueError, match="RTF"):
        extract_spec("spec.docx", rtf)


def test_docx_garbage_friendly_error():
    with pytest.raises(ValueError, match="повреждён или это не DOCX"):
        extract_spec("spec.docx", b"just some text, not a zip")


def test_docx_zip_but_not_docx_friendly_error():
    """ZIP-архив без docx-структуры (PackageNotFoundError) — тоже не 500."""
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("readme.txt", "not a docx")
    with pytest.raises(ValueError, match="Пересохраните"):
        extract_spec("spec.docx", buf.getvalue())


def test_pdf_garbage_friendly_error():
    with pytest.raises(ValueError, match="Не удалось прочитать PDF"):
        extract_spec("spec.pdf", b"garbage not a pdf")


def test_unsupported_extension():
    with pytest.raises(ValueError, match="Неподдерживаемый формат"):
        extract_spec("spec.docx.bak", b"x" * 10)
