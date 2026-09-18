import httpx
import pytest
import respx

from app.config import settings
from app.models import MasterCV

VALID = {"Authorization": f"Bearer {settings.admin_token}"}


@pytest.mark.asyncio
async def test_admin_create_variant_unauthorized_no_token(client):
    # без токена → 401 (а не 422), в едином AppError-формате
    res = await client.post(
        "/api/admin/variants", json={"slug": "x", "title": "t", "content_markdown": "# m"}
    )
    assert res.status_code == 401
    assert res.json()["error"] == "unauthorized"


@pytest.mark.asyncio
async def test_admin_create_variant_unauthorized_wrong_token(client):
    res = await client.post(
        "/api/admin/variants",
        json={"slug": "x", "title": "t", "content_markdown": "# m"},
        headers={"Authorization": "Bearer wrong-token"},
    )
    assert res.status_code == 401


@pytest.mark.asyncio
async def test_admin_create_variant_success(client, session):
    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV", version=1))
    await session.commit()
    res = await client.post(
        "/api/admin/variants",
        headers=VALID,
        json={"slug": "yandex", "title": "Flutter", "content_markdown": "# Yandex CV"},
    )
    assert res.status_code == 201
    body = res.json()
    assert body["slug"] == "yandex"
    assert "id" in body


@pytest.mark.asyncio
async def test_admin_create_variant_conflict_on_duplicate_slug(client, session):
    from app.models import CVVariant, CVVariantStatus

    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV", version=1))
    session.add(
        CVVariant(
            master_cv_id=1, slug="dup", title="t", content_markdown="# m",
            status=CVVariantStatus.draft,
        )
    )
    await session.commit()
    res = await client.post(
        "/api/admin/variants",
        headers=VALID,
        json={"slug": "dup", "title": "t2", "content_markdown": "# m2"},
    )
    assert res.status_code == 409
    assert res.json()["error"] == "conflict"


@pytest.mark.asyncio
async def test_admin_create_link_success(client, session):
    from app.models import CVVariant, CVVariantStatus

    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV", version=1))
    session.add(
        CVVariant(
            master_cv_id=1, slug="sber", title="t", content_markdown="# m",
            status=CVVariantStatus.active,
        )
    )
    await session.commit()
    res = await client.post(
        "/api/admin/links", headers=VALID, json={"cv_variant_slug": "sber", "ttl_days": 14}
    )
    assert res.status_code == 201
    body = res.json()
    assert "code" in body
    assert body["url"].startswith("https://cv.example.com/")
    # код — верхний регистр, 4-6 символов
    assert body["code"].isupper() and 4 <= len(body["code"]) <= 6


@pytest.mark.asyncio
async def test_admin_create_link_unknown_variant_404(client, session):
    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV", version=1))
    await session.commit()
    res = await client.post(
        "/api/admin/links", headers=VALID, json={"cv_variant_slug": "no-such-variant"}
    )
    assert res.status_code == 404


@pytest.mark.asyncio
async def test_admin_create_variant_persists_and_visible_publicly(client, session):
    """Созданный через admin вариант виден через публичный GET /api/variants/{slug}."""
    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV", version=1))
    await session.commit()
    res = await client.post(
        "/api/admin/variants",
        headers=VALID,
        json={
            "slug": "tinkoff",
            "title": "Tinkoff Flutter",
            "content_markdown": "# T CS",
            "status": "active",
        },
    )
    assert res.status_code == 201
    # публичный эндпоинт находит созданный вариант
    pub = await client.get("/api/variants/tinkoff")
    assert pub.status_code == 200
    assert pub.json()["title"] == "Tinkoff Flutter"


@pytest.mark.asyncio
async def test_admin_create_variant_normalizes_slug_to_lowercase(client, session):
    """§4: slug приводится к нижнему регистру."""
    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV", version=1))
    await session.commit()
    res = await client.post(
        "/api/admin/variants",
        headers=VALID,
        json={"slug": "YANDEX", "title": "t", "content_markdown": "# m"},
    )
    assert res.status_code == 201
    assert res.json()["slug"] == "yandex"


# ===== Applications (Отклики) =====


@pytest.mark.asyncio
@respx.mock
async def test_admin_generate_cv(client, session):
    from app.models import MasterCV

    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV\nFlutter", version=1))
    await session.commit()
    sse = (
        'data: {"choices":[{"delta":{"content":"===CV===\\n# CV\\n===COVER===\\n# Cover"}}]}\n\n'
        'data: [DONE]\n\n'
    )
    respx.post(f"{settings.zai_api_base}/chat/completions").mock(
        return_value=httpx.Response(200, text=sse)
    )
    res = await client.post(
        "/api/admin/applications/generate",
        headers=VALID,
        json={"company": "Y", "role": "R", "vacancy_text": "Flutter dev", "selected_projects": []},
    )
    assert res.status_code == 200
    data = res.json()
    assert "# CV" in data["cv_markdown"]
    assert "# Cover" in data["cover_letter"]
    assert "prompt" in data
    assert len(data["prompt"]) > 0  # промпт собран и возвращён


@pytest.mark.asyncio
async def test_admin_create_and_list_application(client, session):
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Acme",
            "role": "Dev",
            "vacancy_text": "Текст вакансии",
            "cv_markdown": "# CV",
            "cover_letter": "# Cover",
            "slug": "acme-dev",
            "generated_prompt": "Ты — помощник... МАСТЕР-CV: # CV\nFlutter",
        },
    )
    assert res.status_code == 201
    app_id = res.json()["id"]

    # список
    res = await client.get("/api/admin/applications", headers=VALID)
    assert res.status_code == 200
    data = res.json()
    assert len(data) == 1
    assert data[0]["company"] == "Acme"
    assert "vacancy_text" not in data[0]  # приватное

    # детальная
    res = await client.get(f"/api/admin/applications/{app_id}", headers=VALID)
    assert res.status_code == 200
    detail = res.json()
    assert detail["vacancy_text"] == "Текст вакансии"
    assert detail["cv_markdown"] == "# CV"
    assert detail["generated_prompt"] == "Ты — помощник... МАСТЕР-CV: # CV\nFlutter"


@pytest.mark.asyncio
async def test_admin_publish_application(client, session):
    # создаём черновик с плейсхолдером ссылки в cover letter
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "X", "role": "R", "vacancy_text": "v",
            "cv_markdown": "# m",
            "cover_letter": "Здравствуйте!\nМоё CV: {CV_LINK}",
            "slug": "x-1",
        },
    )
    app_id = res.json()["id"]
    # публикуем
    res = await client.post(f"/api/admin/applications/{app_id}/publish", headers=VALID)
    assert res.status_code == 200
    data = res.json()
    assert "code" in data
    assert data["url"].startswith("https://")
    short_url = data["url"]
    # плейсхолдер {CV_LINK} должен быть заменён на реальную короткую ссылку
    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert "{CV_LINK}" not in detail["cover_letter"], "плейсхолдер не заменён"
    assert short_url in detail["cover_letter"], "ссылка не вставлена в cover letter"


@pytest.mark.asyncio
async def test_admin_archive_application(client, session):
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Z", "role": "R", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "z-1",
        },
    )
    app_id = res.json()["id"]
    res = await client.post(f"/api/admin/applications/{app_id}/archive", headers=VALID)
    assert res.status_code == 200
    assert res.json()["status"] == "archived"


@pytest.mark.asyncio
async def test_admin_interview_crud(client, session):
    """CRUD интервью: создать → в detail → patch → delete."""
    # создаём отклик
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Int", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "int-1",
        },
    )
    app_id = res.json()["id"]

    # создаём интервью
    res = await client.post(
        f"/api/admin/applications/{app_id}/interviews",
        headers=VALID,
        json={"scheduled_at": "2026-08-01T12:00:00Z", "notes_before": "подготовить стек"},
    )
    assert res.status_code == 201
    iv_id = res.json()["id"]
    assert res.json()["notes_before"] == "подготовить стек"

    # интервью появляется в detail отклика
    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert len(detail["interviews"]) == 1
    assert detail["interviews"][0]["id"] == iv_id

    # редактируем
    res = await client.patch(
        f"/api/admin/interviews/{iv_id}",
        headers=VALID,
        json={"notes_after": "прошло хорошо"},
    )
    assert res.status_code == 200
    assert res.json()["notes_after"] == "прошло хорошо"

    # удаляем
    res = await client.delete(f"/api/admin/interviews/{iv_id}", headers=VALID)
    assert res.status_code == 204

    # больше нет в detail
    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert len(detail["interviews"]) == 0


@pytest.mark.asyncio
async def test_admin_upcoming_interviews(client, session):
    """Дашборд: upcoming возвращает только будущие интервью."""
    # отклик + интервью в будущем
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Up", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "up-1",
        },
    )
    app_id = res.json()["id"]
    await client.post(
        f"/api/admin/applications/{app_id}/interviews",
        headers=VALID,
        json={"scheduled_at": "2099-12-31T10:00:00Z"},
    )

    # upcoming
    res = await client.get("/api/admin/upcoming", headers=VALID)
    assert res.status_code == 200
    data = res.json()
    assert len(data) >= 1
    assert data[0]["application_role"] == "Dev"
    assert data[0]["application_company"] == "Up"


# ===== Тесты артефактов конкурсных откликов =====


@pytest.mark.asyncio
async def test_upload_and_download_artifact(client, session):
    """Загрузка артефакта → публичная ссылка → скачивание."""
    # создаём отклик
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Test", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "art-1",
        },
    )
    app_id = res.json()["id"]

    # загружаем файл
    res = await client.post(
        f"/api/admin/applications/{app_id}/artifacts",
        headers=VALID,
        files={"file": ("clapgo.apk", b"FAKE_APK_BYTES_12345", "application/vnd.android.package-archive")},
    )
    assert res.status_code == 201
    data = res.json()
    assert data["filename"] == "clapgo.apk"
    assert data["size_bytes"] > 0
    assert "/dl/" in data["download_url"]
    code = data["code"]
    artifact_id = data["id"]

    # артефакт виден в деталях отклика
    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert len(detail["artifacts"]) == 1
    assert detail["artifacts"][0]["code"] == code

    # публичное скачивание (без admin-auth)
    res = await client.get(f"/dl/{code}")
    assert res.status_code == 200
    assert b"FAKE_APK" in res.content
    assert "attachment" in res.headers.get("content-disposition", "")

    # удаление
    res = await client.delete(f"/api/admin/artifacts/{artifact_id}", headers=VALID)
    assert res.status_code == 204
    # больше не скачивается
    res = await client.get(f"/dl/{code}")
    assert res.status_code == 404


@pytest.mark.asyncio
async def test_artifact_archived_gone(client, session):
    """Архивированный отклик → 410 Gone при скачивании артефакта."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Gone", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "art-gone",
        },
    )
    app_id = res.json()["id"]

    res = await client.post(
        f"/api/admin/applications/{app_id}/artifacts",
        headers=VALID,
        files={"file": ("test.zip", b"ZIP_DATA", "application/zip")},
    )
    code = res.json()["code"]

    # архивируем
    await client.post(f"/api/admin/applications/{app_id}/archive", headers=VALID)

    # скачивание → 410
    res = await client.get(f"/dl/{code}")
    assert res.status_code == 410


@pytest.mark.asyncio
async def test_artifact_size_limit(client, session, monkeypatch):
    """Превышение лимита размера → 400."""
    from app.config import settings as cfg
    monkeypatch.setattr(cfg, "artifact_max_size_mb", 0)  # 0 MB = 0 bytes max

    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Big", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "art-big",
        },
    )
    app_id = res.json()["id"]

    res = await client.post(
        f"/api/admin/applications/{app_id}/artifacts",
        headers=VALID,
        files={"file": ("big.apk", b"X" * 100, "application/octet-stream")},
    )
    assert res.status_code == 400


# ===== Тесты multi-file spec upload =====


@pytest.mark.asyncio
async def test_upload_spec_docx(client, session):
    """Загрузка DOCX — извлечение текста + таблицы."""
    from docx import Document

    doc = Document()
    doc.add_heading("ТЗ на приложение", level=1)
    doc.add_paragraph("Нужно сделать доставку еды.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Функция"
    table.cell(0, 1).text = "Описание"
    table.cell(1, 0).text = "Корзина"
    table.cell(1, 1).text = "Добавление товаров"

    import io

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = await client.post(
        "/api/admin/applications/upload-spec",
        headers=VALID,
        files={"files": ("tz.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    data = res.json()
    assert "доставку еды" in data["spec_text"]
    assert "Корзина" in data["spec_text"]  # таблица извлечена
    assert len(data["files"]) == 1
    assert data["files"][0]["type"] == "docx"


@pytest.mark.asyncio
async def test_upload_spec_multiple_files(client, session):
    """Загрузка нескольких файлов — тексты объединяются."""
    # DOCX
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("DOCX содержимое")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = await client.post(
        "/api/admin/applications/upload-spec",
        headers=VALID,
        files=[
            ("files", ("doc1.docx", docx_bytes, "application/octet-stream")),
            ("files", ("doc2.docx", docx_bytes, "application/octet-stream")),
        ],
    )
    assert res.status_code == 200
    data = res.json()
    assert "doc1.docx" in data["spec_text"]
    assert "doc2.docx" in data["spec_text"]
    assert "DOCX содержимое" in data["spec_text"]
    assert len(data["files"]) == 2


@pytest.mark.asyncio
async def test_upload_spec_rejects_unsupported(client, session):
    """Неподдерживаемый формат — 400."""
    res = await client.post(
        "/api/admin/applications/upload-spec",
        headers=VALID,
        files={"files": ("file.txt", b"hello", "text/plain")},
    )
    assert res.status_code == 400


# ===== Тесты extra_instruction =====


@pytest.mark.asyncio
async def test_extra_instruction_saved_and_returned(client, session):
    """extra_instruction сохраняется в Application и возвращается в detail."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Test", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "instr-1",
            "extra_instruction": "сделать акцент на backend",
        },
    )
    assert res.status_code == 201
    app_id = res.json()["id"]

    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert detail["extra_instruction"] == "сделать акцент на backend"


@pytest.mark.asyncio
async def test_instructions_endpoint(client, session):
    """GET /instructions — отдаёт только отклики с extra_instruction."""
    # без инструкции — не должен попасть в список
    await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "NoInstr", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "no-instr",
        },
    )
    # с инструкцией
    await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "WithInstr", "role": "Dev", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "with-instr",
            "extra_instruction": "убрать 1С из CV",
        },
    )

    res = await client.get("/api/admin/instructions", headers=VALID)
    assert res.status_code == 200
    data = res.json()
    assert len(data) >= 1
    assert all(d["extra_instruction"] for d in data)
    assert any(d["extra_instruction"] == "убрать 1С из CV" for d in data)


# ===== Тесты platform (kwork) =====


@pytest.mark.asyncio
async def test_platform_saved_and_returned(client, session):
    """platform сохраняется в Application и возвращается в detail."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Kwork", "role": "Web service", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "kwork-1",
            "platform": "kwork",
        },
    )
    assert res.status_code == 201
    app_id = res.json()["id"]

    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert detail["platform"] == "kwork"


@pytest.mark.asyncio
async def test_kwork_active_no_public_link(client, session):
    """kwork-заявка при status=active не получает короткую ссылку,
    CV-вариант остаётся draft — публичная страница отдаёт 404."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Kwork", "role": "Bot", "vacancy_text": "v",
            "cv_markdown": "", "cover_letter": "Отклик", "slug": "kwork-pub-1",
            "platform": "kwork", "status": "active",
        },
    )
    assert res.status_code == 201
    assert res.json().get("url") is None

    app_id = res.json()["id"]
    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert detail["status"] == "active"
    assert detail["short_link_code"] is None

    pub = await client.get("/api/variants/kwork-pub-1")
    assert pub.status_code == 404


@pytest.mark.asyncio
async def test_kwork_publish_marks_sent_without_link(client, session):
    """«Опубликовать» для kwork = отметить отправленным: статус active,
    но без короткой ссылки и без активации варианта."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Kwork", "role": "Parser", "vacancy_text": "v",
            "cv_markdown": "", "cover_letter": "Отклик", "slug": "kwork-pub-2",
            "platform": "kwork", "status": "draft",
        },
    )
    assert res.status_code == 201
    app_id = res.json()["id"]

    pub = await client.post(f"/api/admin/applications/{app_id}/publish", headers=VALID)
    assert pub.status_code == 200
    assert pub.json()["code"] is None
    assert pub.json()["url"] is None

    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert detail["status"] == "active"
    assert detail["short_link_code"] is None
    assert (await client.get("/api/variants/kwork-pub-2")).status_code == 404


@pytest.mark.asyncio
async def test_kwork_generation_no_cv_link(client, session):
    """Kwork промпт не содержит {cv_link} плейсхолдер (запрет ссылок)."""
    from app.llm.generate_prompt import build_generate_prompt
    from app.models import MasterCV

    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV\nDev", version=1))
    await session.commit()

    prompt = await build_generate_prompt(
        session, "# CV\nDev", "Нужен веб-сервис", [], "freelance",
        platform="kwork",
    )
    assert "{cv_link}" not in prompt
    assert "KWORK" in prompt or "kwork" in prompt.lower()
    assert "ЗАПРЕЩЕНЫ" in prompt  # жёсткие правила присутствуют


# ===== Тесты kwork budget (вилка) =====


@pytest.mark.asyncio
async def test_budget_max_saved_and_returned(client, session):
    """budget_max сохраняется и возвращается."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "Kwork", "role": "MVP", "vacancy_text": "v",
            "cv_markdown": "# m", "cover_letter": "", "slug": "kw-budget",
            "platform": "kwork", "budget": "1500", "budget_max": "4500",
        },
    )
    assert res.status_code == 201
    app_id = res.json()["id"]
    detail = (await client.get(f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert detail["budget"] == "1500"
    assert detail["budget_max"] == "4500"


@pytest.mark.asyncio
async def test_kwork_prompt_has_budget_and_estimate(client, session):
    """Kwork промпт с бюджетом: бюджетная вилка + ESTIMATE блок."""
    from app.llm.generate_prompt import build_generate_prompt
    from app.models import MasterCV

    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV\nDev", version=1))
    await session.commit()

    prompt = await build_generate_prompt(
        session, "# CV\nDev", "Нужен MVP", [], "freelance",
        platform="kwork", budget="1500", budget_max="4500",
    )
    assert "Желаемый бюджет: 1500" in prompt
    assert "Допустимый бюджет: 4500" in prompt
    assert "===ESTIMATE===" in prompt
    assert "{cv_link}" not in prompt


# ===== Тесты PDF-экспорта =====


@pytest.mark.asyncio
async def test_pdf_preview_from_markdown(client, session):
    """POST /pdf/preview: PDF из произвольного markdown (несохранённый редактор)."""
    res = await client.post(
        "/api/admin/pdf/preview",
        headers=VALID,
        json={"markdown": "# DevOps Иван\n- Docker, Kubernetes", "title": "DevOps"},
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content[:4] == b"%PDF"
    # имя файла в заголовке — ASCII (latin-1), кириллица вырезается
    assert "CV_DevOps.pdf" in res.headers["content-disposition"]


@pytest.mark.asyncio
async def test_pdf_preview_empty_markdown_400(client, session):
    res = await client.post(
        "/api/admin/pdf/preview", headers=VALID, json={"markdown": "   "}
    )
    assert res.status_code == 400


# ===== Тесты cover_limit и Telegram в промптах =====


@pytest.mark.asyncio
async def test_cover_limit_in_generate_prompt(client, session):
    """cover_limit попадает в собранный промпт генерации."""
    from app.llm.generate_prompt import build_generate_prompt
    from app.models import MasterCV

    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV\nDev", version=1))
    await session.commit()

    prompt = await build_generate_prompt(
        session, "# CV\nDev", "Заказ", [], "freelance", cover_limit=5000,
    )
    assert "НЕ БОЛЕЕ 5000" in prompt


@pytest.mark.asyncio
async def test_no_cover_limit_no_block(client, session):
    """Без cover_limit блок лимита не появляется."""
    from app.llm.generate_prompt import build_generate_prompt
    from app.models import MasterCV

    session.add(MasterCV(id=1, summary="s", contacts={}, full_markdown="# CV\nDev", version=1))
    await session.commit()

    prompt = await build_generate_prompt(
        session, "# CV\nDev", "Заказ", [], "freelance",
    )
    assert "ЛИМИТ ПЛОЩАДКИ" not in prompt


def test_telegram_in_fl_prompts_defaults():
    """Freelance и contest дефолты содержат Telegram-контакт;
    kwork — категорически нет (запрет биржи), vacancy — нет (не FL)."""
    from app.seed_defaults import (
        DEFAULT_PROMPT_GENERATE,
        DEFAULT_PROMPT_GENERATE_CONTEST,
        DEFAULT_PROMPT_GENERATE_FREELANCE,
        DEFAULT_PROMPT_GENERATE_KWORK,
    )

    assert "@vrg18" in DEFAULT_PROMPT_GENERATE_FREELANCE
    assert "@vrg18" in DEFAULT_PROMPT_GENERATE_CONTEST
    assert "@vrg18" not in DEFAULT_PROMPT_GENERATE_KWORK
    assert "@vrg18" not in DEFAULT_PROMPT_GENERATE


# ===== Тесты ESTIMATE: адекватность и риски =====


def test_estimate_risk_fields_in_defaults():
    """Все три генерирующих промпта (freelance/contest/kwork) требуют
    в ===ESTIMATE=== адекватность заказчика и риски."""
    from app.seed_defaults import (
        DEFAULT_PROMPT_GENERATE_CONTEST,
        DEFAULT_PROMPT_GENERATE_FREELANCE,
        DEFAULT_PROMPT_GENERATE_KWORK,
    )

    for name, p in [
        ("freelance", DEFAULT_PROMPT_GENERATE_FREELANCE),
        ("contest", DEFAULT_PROMPT_GENERATE_CONTEST),
        ("kwork", DEFAULT_PROMPT_GENERATE_KWORK),
    ]:
        assert "Адекватность заказчика" in p, name
        assert "Риски заказа" in p, name
        assert "Риски заказчика" in p, name


def test_parse_estimate_with_risk_fields():
    """Парсер принимает расширенный ESTIMATE — все строки попадают в estimate."""
    from app.llm.generate_prompt import parse_generate_response

    raw = """===CV===
# CV
===COVER===
Отклик
===END===
===ESTIMATE===
Оценка стоимости: 40 000-90 000 ₽
Оценка срока: 2-3 недели
Адекватность заказчика: средняя — бюджет ниже объёма
Риски заказа: размытое ТЗ, чужой код
Риски заказчика: возможны бесконечные правки
Комментарий: брать с оговорками"""
    cv, cover, estimate = parse_generate_response(raw)
    assert cv == "# CV"
    assert cover == "Отклик"
    assert estimate is not None
    for line in (
        "Адекватность заказчика",
        "Риски заказа",
        "Риски заказчика",
        "Комментарий",
    ):
        assert line in estimate


# ===== Тесты Negotiation (переговоры с заказчиком) =====


@pytest.mark.asyncio
async def test_negotiation_crud(client, session):
    """Добавление/правка/удаление сообщений переговоров + лента."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "FL", "role": "Интеграция CRM",
            "vacancy_text": "Нужна интеграция CRM с телефонией",
            "cover_letter": "Готов сделать", "cv_markdown": "# CV",
            "slug": "negotiation-crud-1", "kind": "freelance",
        },
    )
    assert res.status_code == 201
    app_id = res.json()["id"]

    # Сообщение заказчика (копипаст с FL)
    r = await client.post(
        f"/api/admin/applications/{app_id}/negotiation",
        headers=VALID,
        json={"role": "customer", "channel": "fl",
              "content": "Здравствуйте! Какой срок и цену можете предложить?"},
    )
    assert r.status_code == 201
    msg_id = r.json()["id"]
    assert r.json()["role"] == "customer"
    assert r.json()["channel"] == "fl"

    # Мой ответ (telegram)
    r = await client.post(
        f"/api/admin/applications/{app_id}/negotiation",
        headers=VALID,
        json={"role": "me", "channel": "telegram", "content": "Срок 2 недели"},
    )
    assert r.status_code == 201

    # Лента: хронология, 2 сообщения
    r = await client.get(f"/api/admin/applications/{app_id}/negotiation", headers=VALID)
    assert r.status_code == 200
    msgs = r.json()
    assert len(msgs) == 2
    assert msgs[0]["role"] == "customer"
    assert msgs[1]["role"] == "me"

    # Правка канала
    r = await client.put(
        f"/api/admin/negotiation/{msg_id}", headers=VALID,
        json={"channel": "telegram"},
    )
    assert r.status_code == 200
    assert r.json()["channel"] == "telegram"

    # Пустое сообщение — 400
    r = await client.post(
        f"/api/admin/applications/{app_id}/negotiation",
        headers=VALID, json={"role": "me", "content": "   "},
    )
    assert r.status_code == 400

    # Удаление
    r = await client.delete(f"/api/admin/negotiation/{msg_id}", headers=VALID)
    assert r.status_code == 204
    r = await client.get(f"/api/admin/applications/{app_id}/negotiation", headers=VALID)
    assert len(r.json()) == 1


@pytest.mark.asyncio
async def test_suggest_reply_validations(client, session):
    """suggest-reply: пустая переписка и «последнее — моё» дают 400."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "FL", "role": "Парсер",
            "vacancy_text": "Нужен парсер", "cover_letter": "Отклик",
            "cv_markdown": "# CV", "slug": "negotiation-suggest-1",
            "kind": "freelance",
        },
    )
    app_id = res.json()["id"]

    # Пустая переписка
    r = await client.post(
        f"/api/admin/applications/{app_id}/suggest-reply", headers=VALID, json={},
    )
    assert r.status_code == 400
    assert "Переписка пуста" in r.json()["message"]

    # Последнее — моё сообщение
    await client.post(
        f"/api/admin/applications/{app_id}/negotiation", headers=VALID,
        json={"role": "customer", "content": "Сколько стоит?"},
    )
    await client.post(
        f"/api/admin/applications/{app_id}/negotiation", headers=VALID,
        json={"role": "me", "content": "50 тысяч"},
    )
    r = await client.post(
        f"/api/admin/applications/{app_id}/suggest-reply", headers=VALID, json={},
    )
    assert r.status_code == 400
    assert "Последнее сообщение" in r.json()["message"]


def test_settings_contains_negotiation_prompt():
    """prompt_negotiation в CONFIG_KEYS и дефолт содержит плейсхолдеры."""
    from app.models.config_text import CONFIG_KEYS
    from app.seed_defaults import DEFAULT_PROMPT_NEGOTIATION

    assert "prompt_negotiation" in CONFIG_KEYS
    for ph in ("{order_text}", "{spec_text}", "{response_text}",
               "{dialog_history}", "{channel}", "{instruction}"):
        assert ph in DEFAULT_PROMPT_NEGOTIATION


# ===== Тесты экспорта заказа (JSON + ZIP) =====


@pytest.mark.asyncio
async def test_export_json_and_zip(client, session):
    """Экспорт содержит все секции; ZIP — структуру папки проекта."""
    import io
    import zipfile as zf

    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "FL", "role": "Бот для такси",
            "vacancy_text": "Нужен бот парсинга заказов",
            "cover_letter": "Сделаю бота за неделю", "cv_markdown": "# CV бота",
            "slug": "export-test-1", "kind": "freelance", "budget": "30000",
            "spec_text": "ТЗ: telebot, 3 команды",
            "estimate": "Оценка стоимости: 25-40k\nРиски заказа: размытое ТЗ",
        },
    )
    app_id = res.json()["id"]

    # Диалог
    await client.post(
        f"/api/admin/applications/{app_id}/negotiation", headers=VALID,
        json={"role": "customer", "channel": "fl", "content": "Когда начнёте?"},
    )
    await client.post(
        f"/api/admin/applications/{app_id}/negotiation", headers=VALID,
        json={"role": "me", "channel": "telegram", "content": "В понедельник"},
    )

    # JSON
    r = await client.get(f"/api/admin/applications/{app_id}/export", headers=VALID)
    assert r.status_code == 200
    data = r.json()
    assert data["meta"]["role"] == "Бот для такси"
    assert data["meta"]["budget"] == "30000"
    assert "бот парсинга" in data["vacancy_text"]
    assert len(data["negotiation"]) == 2
    assert "Риски заказа" in data["estimate"]

    # ZIP
    r = await client.get(f"/api/admin/applications/{app_id}/export.zip", headers=VALID)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/zip"
    assert r.content[:2] == b"PK"
    with zf.ZipFile(io.BytesIO(r.content)) as z:
        names = set(z.namelist())
    assert "AGENTS.md" in names
    assert "order.md" in names
    assert "response.md" in names
    assert "estimate.md" in names
    assert "spec.md" in names
    assert "dialog.md" in names
    with zf.ZipFile(io.BytesIO(r.content)) as z:
        agents = z.read("AGENTS.md").decode()
        assert "Карта файлов" in agents
        dialog = z.read("dialog.md").decode()
        assert "Когда начнёте?" in dialog and "В понедельник" in dialog


@pytest.mark.asyncio
async def test_export_not_found_404(client, session):
    import uuid

    r = await client.get(
        f"/api/admin/applications/{uuid.uuid4()}/export", headers=VALID
    )
    assert r.status_code == 404


# ===== Тесты assistant-треда и черновика ответа =====


@pytest.mark.asyncio
async def test_assistant_thread_and_draft_reply(client, session):
    """История треда, сохранение ответа, очистка; черновик ответа заказчику."""
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "FL", "role": "Доработка магазина",
            "vacancy_text": "Нужна доработка магазина на Opencart",
            "cover_letter": "Сделаю", "cv_markdown": "# CV",
            "slug": "assistant-thread-1", "kind": "freelance",
        },
    )
    app_id = res.json()["id"]

    # Черновик: сохранить и прочитать в detail
    r = await client.put(
        f"/api/admin/applications/{app_id}/draft-reply", headers=VALID,
        json={"draft": "Здравствуйте! Готов начать в понедельник."},
    )
    assert r.status_code == 200
    detail = (await client.get(
        f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert "в понедельник" in detail["draft_reply"]

    # Сохранение ответа ассистента (фронт зовёт после стрима)
    r = await client.post(
        f"/api/admin/applications/{app_id}/assistant-messages", headers=VALID,
        json={"content": "Вот мой совет: сначала запросите ТЗ."},
    )
    assert r.status_code == 201
    assert r.json()["role"] == "assistant"

    # История
    r = await client.get(
        f"/api/admin/applications/{app_id}/assistant-messages", headers=VALID)
    assert len(r.json()) == 1

    # Очистка треда — отклик жив, тред пуст
    r = await client.delete(
        f"/api/admin/applications/{app_id}/assistant-messages", headers=VALID)
    assert r.status_code == 204
    r = await client.get(
        f"/api/admin/applications/{app_id}/assistant-messages", headers=VALID)
    assert r.json() == []
    detail = (await client.get(
        f"/api/admin/applications/{app_id}", headers=VALID)).json()
    assert detail["role"] == "Доработка магазина"  # отклик не удалён


def test_assistant_prompt_markers_in_default():
    """Дефолт ассистента требует ===DRAFT=== и два режима."""
    from app.seed_defaults import DEFAULT_PROMPT_ASSISTANT

    assert "===DRAFT===" in DEFAULT_PROMPT_ASSISTANT
    assert "ДВА РЕЖИМА" in DEFAULT_PROMPT_ASSISTANT


def test_response_edit_prompt_mixed_mode():
    """Дефолт правки отклика — смешанный режим (вопросы без маркеров)."""
    from app.seed_defaults import DEFAULT_PROMPT_RESPONSE_EDIT

    assert "ДВА РЕЖИМА" in DEFAULT_PROMPT_RESPONSE_EDIT
    assert "===COVER===" in DEFAULT_PROMPT_RESPONSE_EDIT
    assert "только его маркер" in DEFAULT_PROMPT_RESPONSE_EDIT


# ===== Тесты staged uploads (файлы до создания заявки) =====


@pytest.mark.asyncio
async def test_staged_uploads_flow(client, session):
    """Файлы любого типа сохраняются сразу; текст best-effort;
    при создании заявки — привязываются и переезжают."""
    import io

    from docx import Document

    d = Document()
    d.add_heading("ТЗ проекта", 0)
    d.add_paragraph("Требование: три экрана")
    buf = io.BytesIO()
    d.save(buf)

    res = await client.post(
        "/api/admin/uploads",
        headers=VALID,
        files=[
            ("files", ("spec.docx", buf.getvalue(), "application/octet-stream")),
            ("files", ("notes.txt", "простой текст заметки".encode(), "text/plain")),
            ("files", ("schema.png", b"\x89PNG fake image", "image/png")),
            ("files", ("broken.docx", b"garbage not a zip", "application/octet-stream")),
        ],
    )
    assert res.status_code == 201
    out = res.json()
    by_name = {r["filename"]: r for r in out}

    assert "Требование: три экрана" in by_name["spec.docx"]["text"]
    assert by_name["spec.docx"]["error"] is None
    assert "заметки" in by_name["notes.txt"]["text"]
    assert by_name["schema.png"]["text"] is None  # бинарный — просто хранится
    assert by_name["schema.png"]["error"] is None
    assert by_name["broken.docx"]["text"] is None
    assert "старый формат" in by_name["broken.docx"]["error"] or "повреждён" in by_name["broken.docx"]["error"]


    ids = [r["id"] for r in out if r["id"]]
    assert len(ids) == 4  # все сохранены (битый docx — тоже, текста нет)

    # Создаём заявку с привязкой
    res = await client.post(
        "/api/admin/applications",
        headers=VALID,
        json={
            "company": "FL", "role": "Схема с файлами",
            "vacancy_text": "заказ", "cover_letter": "ok",
            "cv_markdown": "# CV", "slug": "staged-uploads-1",
            "kind": "freelance", "uploads": ids,
        },
    )
    assert res.status_code == 201
    app_id = res.json()["id"]
    detail = (await client.get(
        f"/api/admin/applications/{app_id}", headers=VALID)).json()
    filenames = {a["filename"] for a in detail["artifacts"]}
    assert {"spec.docx", "notes.txt", "schema.png", "broken.docx"} <= filenames
    for a in detail["artifacts"]:
        assert a["application_id"] == app_id

    # После привязки удалить как staged — нельзя
    aid = detail["artifacts"][0]["id"]
    r = await client.delete(f"/api/admin/uploads/{aid}", headers=VALID)
    assert r.status_code == 400


@pytest.mark.asyncio
async def test_staged_upload_delete_before_create(client, session):
    """Staged-файл можно убрать до создания заявки."""
    res = await client.post(
        "/api/admin/uploads",
        headers=VALID,
        files=[("files", ("tmp.txt", b"temporary", "text/plain"))],
    )
    upload_id = res.json()[0]["id"]
    r = await client.delete(f"/api/admin/uploads/{upload_id}", headers=VALID)
    assert r.status_code == 204
